"""Rasuljon ortoped — Lidlar 2026 jadvalini .xlsx ko'rinishida yaratadi.
Faylni Google Drive'ga yuklab, "Google Sheets orqali ochish" qilsa bo'ldi."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

ROWS = 1000
MANBALAR = ["Instagram", "Telegram", "Facebook", "TikTok", "Tanish orqali", "Qo'ng'iroq", "Boshqa"]
SHIKOYATLAR = ["Bel og'rig'i", "Umurtqa churrasi", "Bo'yin og'rig'i", "Tizza og'rig'i", "Bo'g'im og'rig'i",
               "Skolioz", "Yassi oyoqlik", "Sinish / jarohat", "Artroz", "Boshqa"]
HOLAT_RANG = {
    "Yangi": "CFE2FF", "Javob bermadi": "FDE2E1", "Qayta bog'lanish": "FFF3CD", "O'ylab ko'radi": "FFE8CC",
    "Qabulga yozildi": "D1C4E9", "Keldi": "B7E1CD", "Kelmadi": "F4C7C3", "Rad etdi": "D9D9D9",
}
HOLATLAR = list(HOLAT_RANG)
HEADERS = ["№", "Lid sanasi", "Bemor ismi", "Telefon", "Yoshi", "Manba", "Shikoyat", "Holat",
           "📞 Qachon bog'lanish", "🏥 Qachon keladi", "Soat", "Izoh"]
WIDTHS = [6, 12, 24, 17, 7, 15, 20, 18, 18, 16, 8, 38]

BLUE = "1F4E79"
thin = Side(style="thin", color="D0D7DE")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
fill = lambda c: PatternFill("solid", start_color=c, end_color=c)


def lidlar(ws):
    ws.title = "Lidlar"
    for i, (h, w) in enumerate(zip(HEADERS, WIDTHS), 1):
        c = ws.cell(1, i, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill(BLUE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "D2"

    for r in range(2, ROWS + 2):
        ws.cell(r, 1, f'=IF(C{r}="","",ROW()-1)').font = Font(color="888888")
        for col in range(1, 13):
            c = ws.cell(r, col)
            c.border = border
            c.alignment = Alignment(vertical="center", horizontal="center" if col in (1, 5, 11) else None,
                                    wrap_text=col == 12)
        for col in (2, 9, 10):
            ws.cell(r, col).number_format = "DD.MM.YYYY"
        ws.cell(r, 4).number_format = "@"
        ws.cell(r, 11).number_format = "HH:MM"

    last = ROWS + 1

    def dv(values, col, strict):
        d = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True,
                           showErrorMessage=strict)
        ws.add_data_validation(d)
        d.add(f"{col}2:{col}{last}")

    dv(MANBALAR, "F", False)
    dv(SHIKOYATLAR, "G", False)
    dv(HOLATLAR, "H", True)
    dates = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                           showErrorMessage=True, error="Sanani kiriting, masalan 27.09.2026")
    ws.add_data_validation(dates)
    for col in "BIJ":
        dates.add(f"{col}2:{col}{last}")

    rng = f"A2:L{last}"
    aktiv = 'AND($H2<>"Keldi",$H2<>"Rad etdi",$H2<>"Kelmadi")'
    cf = ws.conditional_formatting
    cf.add(rng, FormulaRule(formula=['AND($J2<>"",$J2=TODAY(),$H2<>"Keldi")'],
                            fill=fill("C8F7C5"), font=Font(bold=True), stopIfTrue=True))
    cf.add(rng, FormulaRule(formula=[f'AND($I2<>"",$I2<TODAY(),{aktiv})'],
                            fill=fill("F8C9C4"), font=Font(color="9C0006"), stopIfTrue=True))
    cf.add(rng, FormulaRule(formula=[f'AND($I2<>"",$I2=TODAY(),{aktiv})'],
                            fill=fill("FFF2A8"), font=Font(bold=True), stopIfTrue=True))
    for h, color in HOLAT_RANG.items():
        cf.add(f"H2:H{last}", FormulaRule(formula=[f'$H2="{h}"'], fill=fill(color)))

    ws["I1"].comment = Comment("🟨 Sariq qator — BUGUN qo'ng'iroq qilish kerak\n"
                               "🟥 Qizil qator — qo'ng'iroq muddati O'TIB KETGAN\n"
                               "🟩 Yashil qator — bemor BUGUN keladi", "Admin", width=320, height=90)
    ws["J1"].comment = Comment("Bemor qabulga keladigan sana. Soatini 'Soat' ustuniga yozing (masalan 14:30).",
                               "Admin", width=300, height=60)
    ws.auto_filter.ref = f"A1:L{last}"


def statistika(ws):
    ws["A1"] = "📊 Statistika"
    ws["A1"].font = Font(size=16, bold=True, color=BLUE)
    for col, w in zip("ABCDE", (30, 10, 3, 24, 10)):
        ws.column_dimensions[col].width = w
    umumiy = [
        ("Jami lidlar", "=COUNTA(Lidlar!C2:C1001)"),
        ("Bugun kelgan lidlar", "=COUNTIF(Lidlar!B2:B1001,TODAY())"),
        ("Shu oy lidlar", '=COUNTIFS(Lidlar!B2:B1001,">="&(EOMONTH(TODAY(),-1)+1),Lidlar!B2:B1001,"<="&EOMONTH(TODAY(),0))'),
        ("Bugun keladigan bemorlar", "=COUNTIF(Lidlar!J2:J1001,TODAY())"),
        ("Bugun qo'ng'iroq kerak", '=COUNTIFS(Lidlar!I2:I1001,"<="&TODAY(),Lidlar!H2:H1001,"<>Keldi",'
                                  'Lidlar!H2:H1001,"<>Rad etdi",Lidlar!H2:H1001,"<>Kelmadi")'),
        ("Konversiya (Keldi / Jami)", '=IFERROR(COUNTIF(Lidlar!H2:H1001,"Keldi")/B3,0)'),
    ]
    for i, (k, f) in enumerate(umumiy, 3):
        ws.cell(i, 1, k).font = Font(bold=True)
        c = ws.cell(i, 2, f)
        c.alignment = Alignment(horizontal="center")
        c.font = Font(size=12)
    ws.cell(3 + len(umumiy) - 1, 2).number_format = "0%"

    def jadval(row, col, title, items, ref):
        for j, t in enumerate((title, "Soni")):
            c = ws.cell(row, col + j, t)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = fill(BLUE)
        for i, it in enumerate(items, row + 1):
            a = ws.cell(i, col, it)
            b = ws.cell(i, col + 1, f"=COUNTIF(Lidlar!{ref}2:{ref}1001,{a.coordinate})")
            b.alignment = Alignment(horizontal="center")
            a.border = b.border = border

    jadval(11, 1, "Holat bo'yicha", HOLATLAR, "H")
    jadval(11, 4, "Manba bo'yicha", MANBALAR, "F")
    jadval(22, 4, "Shikoyat bo'yicha", SHIKOYATLAR, "G")


wb = Workbook()
lidlar(wb.active)
statistika(wb.create_sheet("Statistika"))
wb.save("Rasuljon_ortoped_Lidlar_2026.xlsx")
print("OK: Rasuljon_ortoped_Lidlar_2026.xlsx")
